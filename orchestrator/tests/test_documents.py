"""Turning a document into text, and refusing the ones that cannot become text.

Every refusal is asserted by its message, not just its status code. The message
is the part a user reads and acts on, and the whole reason this feature exists is
that the old frontend refused nothing and said nothing.

Fixtures are built here rather than checked in. A generated PDF is reviewable in
the diff; a binary blob is not.
"""

import io
import re

import docx
import httpx
import pytest
from orchestrator.api.deps import current_user
from orchestrator.config import Settings
from orchestrator.documents.extract import (
    BYTE_LIMIT,
    CHAR_LIMIT,
    extract,
    normalise,
)
from orchestrator.errors import DomainError
from orchestrator.main import create_app, signed_in_for_tests
from reportlab.pdfgen import canvas

SENTENCE = "The system shall record a dispense against a prescription."


def make_pdf(*lines: str, pages: int = 1) -> bytes:
    """A real PDF with a real text layer, built in memory."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for _ in range(pages):
        for index, line in enumerate(lines):
            pdf.drawString(72, 720 - index * 16, line)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_scanned_pdf() -> bytes:
    """A PDF with pages and no text layer, which is what a scan looks like."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.rect(72, 600, 200, 100, fill=0)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_mixed_pdf(*, text_pages: int, blank_pages: int) -> bytes:
    """A PDF with some pages carrying text and the rest carrying nothing.

    Text pages come first, then blank pages, which is what a typed cover page
    in front of scanned body pages looks like.
    """
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for _ in range(text_pages):
        pdf.drawString(72, 720, SENTENCE)
        pdf.showPage()
    for _ in range(blank_pages):
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    buffer = io.BytesIO()
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for cell_index, value in enumerate(row):
                grid.cell(row_index, cell_index).text = value
    document.save(buffer)
    return buffer.getvalue()


def test_reads_a_pdf_and_counts_its_pages() -> None:
    result = extract("SRS.pdf", make_pdf(SENTENCE, pages=3))
    assert SENTENCE in result.text
    assert result.pages == 3
    assert result.format == "pdf"
    assert result.truncated is False


def test_reads_a_docx_including_its_tables() -> None:
    """An SRS states requirements in a table more often than not.

    python-docx keeps table text out of `paragraphs` entirely, so a table shaped
    document extracted as empty until this was handled.
    """
    data = make_docx("Introduction", table=[["FR-1", SENTENCE]])
    result = extract("SRS.docx", data)
    assert "Introduction" in result.text
    assert SENTENCE in result.text
    assert result.pages is None


def test_reads_a_text_file() -> None:
    result = extract("brief.md", b"# Brief\n\n" + SENTENCE.encode())
    assert SENTENCE in result.text
    assert result.format == "md"
    assert result.pages is None


def test_a_short_text_file_is_not_mistaken_for_a_scan() -> None:
    """A tiny brief is a legitimate attachment.

    The no-text-layer check counts characters, so applying it to text formats
    would refuse a one line file with a message about scanned documents.
    """
    result = extract("brief.txt", b"Build a calculator.")
    assert result.text == "Build a calculator."


def test_a_scanned_pdf_is_refused_by_name() -> None:
    with pytest.raises(DomainError) as caught:
        extract("scan.pdf", make_scanned_pdf())
    assert "scanned document" in str(caught.value)
    assert caught.value.status_code == 422


def test_a_thin_one_page_pdf_is_not_mistaken_for_a_scan() -> None:
    """A single requirement is a thin document, not a scan.

    The threshold that caught this was measured across the whole document, so a
    one page PDF carrying one sentence fell under a floor meant to describe a
    page with no text on it at all.
    """
    result = extract("thin.pdf", make_pdf(SENTENCE))
    assert SENTENCE in result.text
    assert result.pages == 1


def test_a_long_scan_with_a_stamped_page_number_is_still_refused() -> None:
    """Per page is what makes this catch it.

    A scanner that stamps "Page 3 of 12" into a text layer leaks enough
    characters across twelve sheets to clear any whole document floor.
    """
    stamped = make_pdf("Page 1 of 12", pages=12)
    with pytest.raises(DomainError) as caught:
        extract("scan.pdf", stamped)
    assert "scanned document" in str(caught.value)


def test_a_partly_scanned_pdf_is_read_and_the_blank_pages_are_reported() -> None:
    """Reading part of a document and saying nothing is the failure this avoids.

    A typed cover page in front of scanned body pages is an ordinary SRS. It has
    real text, so refusing it would be wrong, and it is missing most of its text,
    so accepting it silently would be worse.
    """
    result = extract("mixed.pdf", make_mixed_pdf(text_pages=1, blank_pages=3))
    assert result.pages == 4
    assert result.pages_without_text == 3
    assert SENTENCE in result.text


def test_a_pdf_with_no_readable_page_is_still_refused() -> None:
    with pytest.raises(DomainError) as caught:
        extract("scan.pdf", make_mixed_pdf(text_pages=0, blank_pages=4))
    assert "scanned document" in str(caught.value)


def test_an_empty_file_is_refused() -> None:
    with pytest.raises(DomainError) as caught:
        extract("empty.txt", b"   \n\n  ")
    assert "empty" in str(caught.value)


def test_an_unsupported_type_names_what_is_supported() -> None:
    with pytest.raises(DomainError) as caught:
        extract("data.xlsx", b"anything")
    assert "Cannot read .xlsx files" in str(caught.value)
    assert "PDF, DOCX, TXT or MD" in str(caught.value)


def test_a_file_with_no_extension_is_refused() -> None:
    with pytest.raises(DomainError) as caught:
        extract("README", b"anything")
    assert "no extension" in str(caught.value)


def test_a_renamed_file_is_refused_before_the_parser_sees_it() -> None:
    """Otherwise pypdf raises and the reader gets a traceback, not an instruction."""
    with pytest.raises(DomainError) as caught:
        extract("fake.pdf", b"This is plain text pretending to be a PDF.")
    assert "named .pdf but is not a PDF file" in str(caught.value)


def test_an_oversize_file_names_the_limit() -> None:
    with pytest.raises(DomainError) as caught:
        extract("huge.txt", b"x" * (BYTE_LIMIT + 1))
    assert "The limit is 10 MB" in str(caught.value)
    assert caught.value.status_code == 413


def test_an_oversize_file_is_never_printed_as_the_limit() -> None:
    """Rounding to the nearest whole megabyte refused everything from 10 to
    10.5 MB with "That file is 10 MB. The limit is 10 MB.", a sentence that
    contradicts itself, in the band an oversize file most often lands in."""
    with pytest.raises(DomainError) as caught:
        extract("huge.txt", b"x" * (BYTE_LIMIT + 400 * 1024))
    assert "That file is 10.4 MB. The limit is 10 MB." in str(caught.value)


def test_a_corrupt_pdf_is_refused_by_name() -> None:
    """The parser raised, so nothing can be said about the contents.

    Distinct from the renamed file above, which never reaches pypdf: this one
    begins %PDF- and clears the magic byte check, then falls apart inside it.
    """
    with pytest.raises(DomainError) as caught:
        extract("broken.pdf", b"%PDF-1.7\nno xref, no trailer, no pages\n")
    assert "could not be read. It may be corrupt" in str(caught.value)
    assert caught.value.status_code == 422


def test_a_corrupt_docx_is_refused_by_name() -> None:
    with pytest.raises(DomainError) as caught:
        extract("broken.docx", b"PK\x03\x04" + b"not the rest of a zip")
    assert "could not be read. It may be corrupt" in str(caught.value)
    assert caught.value.status_code == 422


def test_truncation_is_reported_rather_than_silent() -> None:
    body = ("The system shall do a thing. " * 4000).encode()
    result = extract("long.txt", body)
    assert result.truncated is True
    assert result.chars == CHAR_LIMIT
    assert result.chars_available > CHAR_LIMIT


def test_normalise_removes_extraction_noise_without_changing_words() -> None:
    """PDF extraction emits blank lines and trailing spaces nobody typed.

    That noise reaches split_sentences() in C1, which sets the extraction budget,
    so it is removed here rather than tolerated downstream.
    """
    assert normalise("a  \n\n\n\n b \r\n c \n\n") == "a\n\n b\n c"


#: Asserted identically by the frontend in entities/document/extract.test.ts.
#: The two normalise implementations exist because the browser reads text files
#: itself; this vector is what stops them drifting apart.
SHARED_VECTOR_IN = "Requirement one.  \r\n\r\n\r\n\r\nRequirement two.\t\n"
SHARED_VECTOR_OUT = "Requirement one.\n\nRequirement two."

#: The other half of that vector, also asserted in extract.test.ts.
#:
#: The normalisation vector is ASCII, which is the one input class where
#: counting characters cannot diverge: Python counts code points and JavaScript
#: counts UTF-16 code units, and the two agree until a character needs two of
#: them. One leading "a" then astral emoji puts every emoji on an odd code unit
#: boundary, so the browser's cut used to land between the halves of one.
ASTRAL_VECTOR = "a" + "\U0001f600" * CHAR_LIMIT


def test_shared_normalisation_vector() -> None:
    assert normalise(SHARED_VECTOR_IN) == SHARED_VECTOR_OUT


def test_shared_astral_truncation_vector() -> None:
    """One emoji is one character, on this side and in the browser."""
    result = extract("emoji.md", ASTRAL_VECTOR.encode())
    assert result.chars_available == CHAR_LIMIT + 1
    assert result.chars == CHAR_LIMIT
    assert result.truncated is True
    assert result.text.endswith("\U0001f600")


def test_payload_is_camel_case() -> None:
    payload = extract("brief.txt", SENTENCE.encode()).as_payload()
    assert set(payload) == {
        "name",
        "format",
        "bytes",
        "pages",
        "pagesWithoutText",
        "words",
        "chars",
        "charsAvailable",
        "truncated",
        "text",
    }


def test_word_count_matches_the_text_returned() -> None:
    result = extract("brief.txt", SENTENCE.encode())
    assert result.words == len(re.findall(r"\S+", result.text))


@pytest.fixture
def client():
    app = create_app(Settings(database_url="postgresql://sdlc:sdlc@localhost:5432/sdlc"))
    # No lifespan and no database here: extraction reads nothing stored, so the
    # request is signed in by override rather than by a session it would look up.
    app.dependency_overrides[current_user] = signed_in_for_tests
    transport = httpx.ASGITransport(app=app)
    return httpx.AsyncClient(transport=transport, base_url="http://test")


async def test_upload_returns_the_text(client: httpx.AsyncClient) -> None:
    async with client:
        response = await client.post(
            "/documents/extract",
            files={"file": ("SRS.pdf", make_pdf(SENTENCE), "application/pdf")},
        )
    assert response.status_code == 200
    body = response.json()
    assert SENTENCE in body["text"]
    assert body["pages"] == 1
    assert body["charsAvailable"] == body["chars"]


async def test_a_refusal_reaches_the_browser_as_a_reason(client: httpx.AsyncClient) -> None:
    """`lib/http.ts` surfaces the parsed body, so the reason has to be in it."""
    async with client:
        response = await client.post(
            "/documents/extract",
            files={"file": ("scan.pdf", make_scanned_pdf(), "application/pdf")},
        )
    assert response.status_code == 422
    assert "scanned document" in response.json()["error"]


async def test_an_oversize_upload_is_refused(client: httpx.AsyncClient) -> None:
    async with client:
        response = await client.post(
            "/documents/extract",
            files={"file": ("huge.txt", b"x" * (BYTE_LIMIT + 1024), "text/plain")},
        )
    assert response.status_code == 413
    assert "10 MB" in response.json()["error"]
